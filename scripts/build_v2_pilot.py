"""Rebuild the isolated V2 pilot; no retrieval, judgments, model or index loading.

Run from the repository with PYTHONPATH=src using the ingestion environment.
The tokenizer subprocess uses only the installed Rust tokenizers package and
the checksum-verified local tokenizer.json, never transformers/torch weights.
"""
import argparse
from collections import Counter
from dataclasses import asdict
from datetime import datetime, timezone
import csv
import io
import json
import math
from pathlib import Path
import re
import statistics
import subprocess
import sys
import tomllib

from openpyxl import load_workbook
from wo_retrieval.pipeline import (build_chunks, ingest, json_text, jsonl_text, load_config,
                                   safe_write, sha256, verify_canonical, write_chunks, write_ingestion)
from wo_retrieval.retrieval.text import serialize


STATUS = "ai_corrected_pilot_not_human_validated"
SNAPSHOT_ID = "WO-CORPUS-V2-AI-CORRECTED-PILOT-001"


def digest(value):
    import hashlib
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def text_hash(value):
    import hashlib
    return hashlib.sha256(value.encode()).hexdigest()


def rows(path):
    return [json.loads(line) for line in path.read_text(encoding="utf8").splitlines()]


def logical_locator(chunk):
    # Compare underlying source blocks independently of chunk IDs and text offsets.
    refs = chunk["source_locator"]["blocks"]
    return [{k: v for k, v in ref.items() if k not in {"serialized_char_start", "serialized_char_end"}} for ref in refs]


def compare_chunk(old, candidates, policy):
    same = [new for new in candidates if logical_locator(new) == logical_locator(old)]
    exact = [new for new in same if new["content_hash"] == old["content_hash"] and new["text"] == old["text"]]
    selected = exact if exact else same
    if not selected:
        return "MISSING_OR_INVALID", None
    if len(selected) != 1:
        return "CHUNK_MAPPING_AMBIGUOUS", None
    new = selected[0]
    unchanged = (new["text"] == old["text"] and new["content_hash"] == old["content_hash"]
                 and serialize(new, policy)[0] == serialize(old, policy)[0])
    return ("EXACT_CONTENT_MATCH_CANDIDATE" if unchanged else "CHANGED_CONTENT_REQUIRES_REJUDGMENT"), new


def quality(documents, chunks):
    issues, covered, operational = [], 0, 0
    orphan = []
    by_doc = {d.document_id: d for d in documents}
    for c in chunks:
        d = by_doc[c.document_id]
        block_map = {b.block_id: b for b in d.content_blocks}
        if not c.text.strip() or c.title != d.title or not c.heading:
            issues.append({"chunk_id": c.chunk_id, "problem": "empty/title/heading"})
        if c.content_hash != text_hash(c.text) or c.token_count != math.ceil(len(c.text)/4) or c.token_count > 500:
            issues.append({"chunk_id": c.chunk_id, "problem": "hash/proxy_token_bound"})
        for ref in c.source_locator["blocks"] + c.source_locator.get("overlap_blocks", []):
            b = block_map.get(ref["block_id"])
            if b is None or any(ref.get(k) != v for k, v in b.source_locator.items()):
                issues.append({"chunk_id": c.chunk_id, "problem": "invalid_source_locator"})
        body = c.text.split("\n\n", 1)[-1].strip()
        flags = []
        if len(body) < 80 and not re.search(r"[.!?]", body): flags.append("very_short_non_sentence")
        if re.fullmatch(r"(?:lihat|rujuk|see)\b.{0,90}", body, re.I): flags.append("cross_reference_only")
        if re.fullmatch(r"[^\n:]{1,80}:\s*", body): flags.append("label_without_value")
        if c.metadata["format"] == "XLSX" and c.source_locator.get("sheet") != "Panduan" and "Header:" not in c.text:
            flags.append("table_without_header")
        if flags: orphan.append({"chunk_id": c.chunk_id, "flags": flags, "text": c.text})
    for d in documents:
        for b in d.content_blocks:
            if not b.retrievable or b.kind in {"title", "heading", "table_header"}: continue
            operational += 1
            original = f"Row: {b.source_locator['row_start']}\n{b.text}" if b.kind == "table_row" else b.text
            positions = set()
            for c in chunks:
                if c.document_id != d.document_id: continue
                for ref in c.source_locator["blocks"]:
                    if ref["block_id"] != b.block_id: continue
                    start, end = ref["serialized_char_start"], ref["serialized_char_end"]
                    if not (0 <= start < end <= len(original)) or original[start:end] not in c.text:
                        issues.append({"chunk_id": c.chunk_id, "problem": "invalid_text_span"})
                    positions.update(range(start, end))
            if all(i in positions or char.isspace() for i, char in enumerate(original)): covered += 1
            else: issues.append({"document_id": d.document_id, "block_id": b.block_id, "problem": "uncovered"})
    duplicate = len(chunks)-len({c.content_hash for c in chunks})
    assert len({c.chunk_id for c in chunks}) == len(chunks)
    assert set(c.document_id for c in chunks) == set(by_doc)
    return {"status": "PASS" if not issues and not duplicate and not orphan else "FAIL",
            "errors": issues, "orphan_candidates": orphan, "orphan_policy": "Deterministic limited fragment heuristics; short chunks not automatically bad; not semantic validation",
            "duplicate_content_hashes": duplicate, "operational_blocks": operational, "covered_blocks": covered,
            "missing_documents": [], "empty_chunks": sum(not c.text.strip() for c in chunks)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--tokenizer-python", type=Path, required=True)
    args = parser.parse_args()
    root = args.root.resolve()
    cfg = load_config(root, root/"experiments/configs/pilot_v2.toml")
    v1cfg = load_config(root, root/"experiments/configs/pilot.toml")
    assert cfg["chunking"] == v1cfg["chunking"]
    base = root/"data/versions/v2_ai_corrected"
    if base.exists(): raise ValueError("V2 output already exists; refuse to overwrite a pilot snapshot")
    old_snapshot = json.loads((root/"data/validation/pilot_snapshot.json").read_text())
    for name, h in old_snapshot["artifact_sha256"].items(): assert sha256(root/name)==h
    assert digest(old_snapshot["artifact_sha256"]) == old_snapshot["aggregate_fingerprint"]["value"]
    workbook = root/"data/validation/ai_corpus_correction_ledger.xlsx"
    wb = load_workbook(workbook, read_only=True, data_only=True)
    sheets = {s.title: [dict(zip(next(s.values), row)) for row in list(s.values)[1:]] for s in wb}
    wb.close()
    file_records = sheets["Files"]
    assert len(file_records)==50 and sum(f["change_status"]=="BYTE_IDENTICAL" for f in file_records)==48
    assert {f["document_id"] for f in file_records if f["change_status"]=="CHANGED"} == {"WO-DOC-024", "WO-DOC-027"}
    for f in file_records:
        assert sha256(root/f["source_path"]) == f["original_sha256"]
        assert sha256(root/f["candidate_path"]) == f["candidate_sha256"]
    unresolved = [{**x, "pilot_blocker": False,
                   "pilot_disposition": "Retain unresolved with synthetic/non-human-validated disclosure; not authorized operational guidance"}
                  for x in sheets["Corrections"] if x["action"]=="UNRESOLVED_NO_PATCH"]
    assert {x["issue_id"] for x in unresolved}=={"C06","C07","C08","C09","C10","C12"}
    manifest = json.loads((root/"data/source_corpus/corpus_manifest.json").read_text())
    for record in manifest:
        path = root/cfg["source_documents"]/record["filename"]
        record["parent_file_sha256"] = record["file_sha256"]
        record["parent_dataset_version"] = record["dataset_version"]
        record.update(file_sha256=sha256(path), dataset_version=cfg["corpus_version"], validation_status=STATUS)
        record["correction_ledger_path"] = workbook.relative_to(root).as_posix()
        record["source_path"] = path.relative_to(root).as_posix()
        if record["document_id"] in {"WO-DOC-024", "WO-DOC-027"}:
            # The two edited files are prose-only DOCX; keep descriptive metadata honest.
            from docx import Document
            text = "\n".join(p.text for p in Document(path).paragraphs)
            record["actual_word_count"] = len(text.split())
            record["actual_character_count"] = len(text)
            record["candidate_count_method"] = "DOCX paragraph text joined with LF; words split on whitespace"
    safe_write(root, root/cfg["manifest"], json_text(manifest))
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=["relative_path","filename","file_extension","file_size_bytes","sha256","file_role","validation_status"])
    writer.writeheader()
    for m in manifest:
        path = root/cfg["source_documents"]/m["filename"]
        writer.writerow(dict(relative_path=f"documents/{m['filename']}", filename=m["filename"], file_extension=path.suffix,
                             file_size_bytes=path.stat().st_size, sha256=m["file_sha256"], file_role="corpus_document", validation_status=STATUS))
    safe_write(root, root/cfg["inventory"], stream.getvalue())
    documents, ingestion, source_audit = ingest(root, cfg)
    if ingestion["failed_documents"]: raise ValueError(ingestion["failures"])
    verify_canonical(root, cfg, documents)
    chunks, chunk_report = build_chunks(documents, cfg)
    repeated, _, _ = ingest(root, cfg)
    repeated_chunks, _ = build_chunks(repeated, cfg)
    assert jsonl_text(documents)==jsonl_text(repeated)
    assert jsonl_text(chunks)==jsonl_text(repeated_chunks)
    write_ingestion(root, cfg, documents, ingestion, source_audit)
    write_chunks(root, cfg, documents, chunks, chunk_report)
    report_dir = base/"reports"
    def output(name, value): safe_write(root, report_dir/name, json_text(value))
    qa = quality(documents, chunks)
    output("quality_audit.json", qa)
    if qa["status"] != "PASS": raise ValueError("Chunk quality gate failed; snapshot not published")
    old_docs = rows(root/"data/processed/documents.jsonl")
    old_chunks = rows(root/"data/chunks/chunks.jsonl")
    new_chunks = [asdict(c) for c in chunks]
    protocol = tomllib.loads((root/"experiments/configs/retrieval_protocol_v1.toml").read_text(encoding="utf8"))
    policy = protocol["serialization"]
    serialized = [{"chunk_id": c["chunk_id"], "text": serialize(c, policy)[0], "text_only": c["text"]} for c in new_chunks]
    safe_write(root, report_dir/"serialized_inputs.jsonl", "".join(json.dumps(x, ensure_ascii=False, sort_keys=True)+"\n" for x in serialized))
    lock = json.loads((root/"experiments/configs/ml_environment.lock.json").read_text())
    model = lock["model"]
    assert model["repo_id"]==protocol["R3"]["model_id"] and model["revision"]==protocol["R3"]["revision"]
    tokenizer_path = Path(model["snapshot_path"])/"tokenizer.json"
    expected = next(f["sha256"] for f in model["files"] if f["file"]=="tokenizer.json")
    assert sha256(tokenizer_path)==expected
    code = """import sys,json
from pathlib import Path
from tokenizers import Tokenizer
tokenizer=Tokenizer.from_file(sys.argv[1]);tokenizer.no_truncation();tokenizer.no_padding()
records=[]
for line in Path(sys.argv[2]).read_text(encoding='utf8').splitlines():
 r=json.loads(line);a=tokenizer.encode(r['text_only'],add_special_tokens=True);b=tokenizer.encode(r['text'],add_special_tokens=True)
 assert a.ids[0]==b.ids[0]==0 and a.ids[-1]==b.ids[-1]==2
 records.append(dict(chunk_id=r['chunk_id'],text_only_tokens=len(a.ids),serialized_tokens=len(b.ids),serialized_text_sha256=__import__('hashlib').sha256(r['text'].encode()).hexdigest()))
Path(sys.argv[3]).write_text(json.dumps(records,ensure_ascii=False,sort_keys=True,indent=2)+'\\n',encoding='utf8')
assert 'torch' not in sys.modules and 'transformers' not in sys.modules
"""
    import os
    env = {**os.environ, "PYTHONDONTWRITEBYTECODE":"1", "HF_HUB_OFFLINE":"1", "TRANSFORMERS_OFFLINE":"1",
           "HF_HOME":"D:/ModelCache/huggingface", "HF_HUB_CACHE":"D:/ModelCache/huggingface/hub", "TORCH_HOME":"D:/ModelCache/torch",
           "PIP_CACHE_DIR":"D:/ModelCache/pip", "TMP":"D:/ModelCache/tmp", "TEMP":"D:/ModelCache/tmp", "TOKENIZERS_PARALLELISM":"false"}
    subprocess.run([str(args.tokenizer_python.resolve()),"-B","-c",code,str(tokenizer_path),str(report_dir/"serialized_inputs.jsonl"),str(report_dir/"tokenizer_lengths.json")],env=env,check=True,timeout=120)
    lengths=json.loads((report_dir/"tokenizer_lengths.json").read_text())
    def stats(key):
        values=sorted(r[key] for r in lengths)
        return dict(min=min(values),median=statistics.median(values),mean=statistics.mean(values),p95_nearest_rank=values[math.ceil(.95*len(values))-1],max=max(values),overflow_at_640=sum(v>640 for v in values))
    token_audit={"model_id":model["repo_id"],"revision":model["revision"],"tokenizer_file_sha256":expected,
                 "method":"Local tokenizers.Tokenizer.from_file; no truncation/padding; special tokens BOS0/EOS2 included; no neural model/torch/transformers import",
                 "text_only":stats("text_only_tokens"),"serialized":stats("serialized_tokens"),"chunks":len(chunks),"max_length":640}
    output("tokenizer_audit.json",token_audit)
    if token_audit["serialized"]["overflow_at_640"]: raise ValueError("Tokenizer length gate failed; snapshot not published")
    mapping=[]
    for old in old_chunks:
        classification,new=compare_chunk(old,[c for c in new_chunks if c["document_id"]==old["document_id"]],policy)
        mapping.append(dict(v1_chunk_id=old["chunk_id"],v2_chunk_id=new["chunk_id"] if new else None,document_id=old["document_id"],classification=classification,
                            v1_content_hash=old["content_hash"],v2_content_hash=new["content_hash"] if new else None,
                            v1_locator=old["source_locator"],v2_locator=new["source_locator"] if new else None,
                            v1_serialized_hash=text_hash(serialize(old,policy)[0]),v2_serialized_hash=text_hash(serialize(new,policy)[0]) if new else None,
                            evidence="Unique document+underlying source-block locator; exact text/content and serialized hashes checked; IDs alone not used"))
    output("v1_v2_chunk_mapping.json",{"status":"AUDIT_ONLY_NOT_EMBEDDING_REUSE_AUTHORIZATION","rows":mapping,"counts":dict(Counter(m["classification"] for m in mapping))})
    mapped={m["v1_chunk_id"]:m for m in mapping}
    old_by_id={c["chunk_id"]:c for c in old_chunks}
    query_rows=rows(root/"data/benchmark/development_queries_draft.jsonl")
    query_map={q["query_id"]:q for q in query_rows}
    assert len(query_rows)==len(query_map)==18
    judgments=rows(root/"data/benchmark/development_qrels_ai_silver_v1.jsonl")
    assert len({(j["query_id"],j["chunk_id"]) for j in judgments})==len(judgments)==684
    impacts=[]
    for j in judgments:
        m=mapped.get(j["chunk_id"]);c=old_by_id.get(j["chunk_id"]);q=query_map.get(j["query_id"])
        valid=bool(m and c and q and c["document_id"]==j["document_id"] and j["chunk_text_sha256"]==text_hash(c["text"])
                   and j["query_text_sha256"]==text_hash(q["text"]) and j["query_text"]==q["text"]
                   and j["snapshot_id"]==old_snapshot["snapshot_id"] and j["source_locator"]==c["source_locator"]
                   and j["source_file"]==c["source_file"])
        impacts.append(dict(query_id=j["query_id"],v1_chunk_id=j["chunk_id"],candidate_v2_chunk_id=m["v2_chunk_id"] if m else None,
                            classification=m["classification"] if valid else "MISSING_OR_INVALID",automatic_transfer=False))
    assert len(impacts)==684
    output("silver_qrels_impact.json",{"status":"AUDIT_ONLY_NO_V2_QRELS_OR_GRADES","counts":dict(Counter(i["classification"] for i in impacts)),"rows":impacts,
                                       "v2_query_chunk_pairs":len(query_rows)*len(chunks),"v2_accepted_judgments":0,"v2_unjudged_pairs":len(query_rows)*len(chunks)})
    new_doc_map={d.document_id:d for d in documents};old_doc_map={d["document_id"]:d for d in old_docs}
    provenance=rows(root/"data/benchmark/development_query_provenance.jsonl")
    query_impacts=[]
    for p in provenance:
        evidence=[]
        for ref in p["source_evidence_refs"]:
            did,bid=ref["document_id"],ref["canonical_block_id"]
            old=next((b for b in old_doc_map[did]["content_blocks"] if b["block_id"]==bid),None)
            new=next((b for b in new_doc_map[did].content_blocks if b.block_id==bid),None)
            assert old and new and old["source_locator"]==ref["source_locator"] and old["heading"]==ref["heading"]
            assert ref["document_title"]==old_doc_map[did]["title"] and ref["source_file"]==old_doc_map[did]["source_file"]
            assert ref["source_format"]==old_doc_map[did]["metadata"]["format"]
            associated=[m for m in mapping if m["document_id"]==did and any(r["block_id"]==bid for r in m["v1_locator"]["blocks"])]
            evidence.append(dict(document_id=did,block_id=bid,locator_valid=new.source_locator==ref["source_locator"],heading_valid=new.heading==ref["heading"],
                                 text_unchanged=old["text"]==new.text,source_hash_unchanged=old_doc_map[did]["metadata"]["file_sha256"]==new_doc_map[did].metadata["file_sha256"],
                                 v1_block_text_sha256=text_hash(old["text"]),v2_block_text_sha256=text_hash(new.text),
                                 containing_chunk_mapping=[dict(v1_chunk_id=m["v1_chunk_id"],v2_chunk_id=m["v2_chunk_id"],classification=m["classification"]) for m in associated]))
        query_impacts.append(dict(query_id=p["query_id"],query_text_unchanged=True,evidence=evidence,
                                 answerability_status="REVIEW_CHANGED_EVIDENCE" if any(not e["text_unchanged"] for e in evidence) else "NO_DIRECT_EVIDENCE_TEXT_CHANGE_DETECTED_NOT_REVALIDATED",
                                 required_action="Rebind all pointers to V2 source paths/snapshot only in a separately approved query-version task; no query or provenance edited"))
    assert len(query_impacts)==18 and sum(len(q["evidence"]) for q in query_impacts)==42
    assert all(e["locator_valid"] and e["heading_valid"] for q in query_impacts for e in q["evidence"])
    output("query_provenance_impact.json",{"status":"AUDIT_ONLY_NOT_QUERY_VALIDATION","rows":query_impacts})
    comparison=[]
    for d in documents:
        old=old_doc_map[d.document_id]
        old_body=[{k:b[k] for k in ('block_id','kind','text','heading','source_locator','retrievable')} for b in old["content_blocks"]]
        new_body=[{k:asdict(b)[k] for k in ('block_id','kind','text','heading','source_locator','retrievable')} for b in d.content_blocks]
        oc=[c for c in old_chunks if c["document_id"]==d.document_id];nc=[c for c in new_chunks if c["document_id"]==d.document_id]
        comparison.append(dict(document_id=d.document_id,extracted_blocks_unchanged=old_body==new_body,
                               operational_text_unchanged=[b["text"] for b in old["content_blocks"] if b["retrievable"]]==[b.text for b in d.content_blocks if b.retrievable],
                               v1_chunks=len(oc),v2_chunks=len(nc),full_span_locators_unchanged=[c["source_locator"] for c in oc]==[c["source_locator"] for c in nc],
                               block_grouping_unchanged=[logical_locator(c) for c in oc]==[logical_locator(c) for c in nc],
                               serialization_unchanged=[serialize(c,policy)[0] for c in oc]==[serialize(c,policy)[0] for c in nc],
                               xlsx_context_unchanged=(all(a["metadata"].get("guide_context")==b["metadata"].get("guide_context") and a["source_locator"]==b["source_locator"] and a["text"]==b["text"] for a,b in zip(oc,nc)) and len(oc)==len(nc)) if d.metadata["format"]=="XLSX" else None))
    old_ids={c["chunk_id"] for c in old_chunks};new_ids={c["chunk_id"] for c in new_chunks}
    diff={"v1_chunks":len(old_chunks),"v2_chunks":len(chunks),"documents":comparison,"mapping_counts":dict(Counter(m["classification"] for m in mapping)),
          "added_chunk_ids":sorted(new_ids-old_ids),"removed_chunk_ids":sorted(old_ids-new_ids),
          "modified_chunk_ids":[m["v1_chunk_id"] for m in mapping if m["classification"]=="CHANGED_CONTENT_REQUIRES_REJUDGMENT"],
          "deterministic_rebuild_documents_sha256":text_hash(jsonl_text(repeated)),"deterministic_rebuild_chunks_sha256":text_hash(jsonl_text(repeated_chunks))}
    output("v1_v2_diff.json",diff)
    output("outstanding_issues.json",{"status":STATUS,"issues":unresolved,"critical_pilot_blockers":[]})
    output("serialization_audit.json",{"policy_id":policy["policy_id"],"policy_sha256":digest(policy),"protocol_sha256":sha256(root/"experiments/configs/retrieval_protocol_v1.toml"),
                                        "serialized_inputs_sha256":sha256(report_dir/"serialized_inputs.jsonl"),"changed_inputs":diff["modified_chunk_ids"],
                                        "xlsx_context_parity":all(d["xlsx_context_unchanged"] for d in comparison if d["xlsx_context_unchanged"] is not None),
                                        "metadata_allowlist_only":True,"neural_model_loaded":False})
    for name,h in old_snapshot["artifact_sha256"].items(): assert sha256(root/name)==h
    for f in file_records: assert sha256(root/f["candidate_path"])==f["candidate_sha256"]
    artifacts={p.relative_to(root).as_posix():sha256(p) for p in sorted(base.rglob('*')) if p.is_file()}
    for name in ['experiments/configs/pilot_v2.toml','experiments/configs/pilot.toml','experiments/configs/retrieval_protocol_v1.toml',
                 'data/validation/ai_corpus_correction_ledger.xlsx','experiments/reports/milestone_9a5_corpus_correction.md','scripts/build_v2_pilot.py']:
        artifacts[name]=sha256(root/name)
    artifacts.update({p.relative_to(root).as_posix():sha256(p) for p in sorted((root/'src/wo_retrieval').rglob('*.py'))})
    snapshot={"snapshot_id":SNAPSHOT_ID,"corpus_version":cfg['corpus_version'],"validation_status":STATUS,
              "created_at_utc":datetime.now(timezone.utc).isoformat(),"purpose":"Preliminary pilot retrieval benchmarking inputs only; not final research dataset or human/expert validation",
              "parent_snapshot_id":old_snapshot['snapshot_id'],"parent_snapshot_sha256":sha256(root/'data/validation/pilot_snapshot.json'),
              "parent_fingerprint":old_snapshot['aggregate_fingerprint']['value'],"source_document_count":len(documents),"canonical_document_count":len(documents),"chunk_count":len(chunks),
              "source_format_distribution":ingestion['formats'],"chunk_format_distribution":chunk_report['chunks_by_format'],"chunking":cfg['chunking'],
              "configuration_provenance":{"config_path":"experiments/configs/pilot_v2.toml","parent_config_path":"experiments/configs/pilot.toml","chunking_parameters_identical":True,
                                          "retrieval_protocol_still_bound_to_v1":True,"retrieval_parameters_unchanged":True,"environment":ingestion['environment']},
              "artifact_sha256":artifacts,"source_file_sha256":{f['candidate_path']:f['candidate_sha256'] for f in file_records},
              "aggregate_fingerprint":{"algorithm":"sha256","serialization":"UTF8 sorted-key compact JSON of artifact_sha256; no trailing newline; source manifest anchors all source hashes; snapshot timestamp excluded", "value":digest(artifacts)},
              "correction_ledger_path":"data/validation/ai_corpus_correction_ledger.xlsx","outstanding_issue_ids":[i['issue_id'] for i in unresolved],
              "technical_validation":{"chunk_quality":qa['status'],"deterministic_rebuild":"PASS","tokenizer_overflow_640":0,"DOCX_visual_QA":"NOT_TESTED","human_validation":"NOT_PERFORMED"},
              "qrels":{"v2_created":False,"existing_v1_judgments":684,"v2_query_chunk_pairs":len(query_rows)*len(chunks),"accepted_v2_judgments":0},
              "indices":{"v2_rebuilt":False,"v1_unchanged":True,"existing_indices_not_compatible_with_v2_snapshot":True}}
    safe_write(root,base/'validation/pilot_snapshot.json',json_text(snapshot))
    print(json_text({"snapshot_id":SNAPSHOT_ID,"fingerprint":snapshot['aggregate_fingerprint']['value'],"chunks":len(chunks),
                     "chunk_formats":chunk_report['chunks_by_format'],"mapping":diff['mapping_counts'],"qrels":dict(Counter(x['classification'] for x in impacts)),"tokenizer":token_audit}))


if __name__ == '__main__':
    main()
