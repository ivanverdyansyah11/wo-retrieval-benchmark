import hashlib
import json
import shutil
import tempfile
import unittest
from dataclasses import asdict
from pathlib import Path

from docx import Document as WriteDocx
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from wo_retrieval.ingestion import ParseError, parse
from wo_retrieval.pipeline import (audit, build_chunks, ingest, jsonl_text, load_config,
                                   read_documents, safe_write, snapshot, verify_canonical)
from wo_retrieval.preprocessing.chunking import chunk_document, token_count
from wo_retrieval.preprocessing.cleaning import clean, normalize
from wo_retrieval.schemas import Block, Document
from wo_retrieval.reporting import markdown_report


ROOT = Path(__file__).resolve().parents[1]
PILOT = {"target_tokens": 500, "overlap_tokens": 75,
         "token_count_method": "approx_unicode_chars_div4_v1"}


def document(blocks, title="Fixture"):
    return Document("WO-DOC-999", title, "test", "fixture", "fixture.txt",
                    "researcher_composed_operational_draft", "1.0", "v1-draft",
                    "pending_practitioner_validation", {"format": "TXT", "file_sha256": "fixture"}, blocks)


def pdf_fixture(path, empty=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=500, height=500)
    if not empty:
        regular = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        bold = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica-Bold")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(regular), NameObject("/F2"): writer._add_object(bold)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F2 12 Tf 30 450 Td (Procedure) Tj ET BT /F1 10 Tf 30 420 Td (Do not change STAFF-01 or 16:30.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_docx_structure_and_order(self):
        path = self.base / "sample.docx"
        d = WriteDocx()
        d.add_heading("Fixture", 0)
        d.add_heading("Procedure", 1)
        d.add_paragraph("Jangan mengubah STAFF-01 pada 16:30.")
        d.add_paragraph("Periksa akses", style="List Bullet")
        d.add_paragraph("Konfirmasi", style="List Number")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "PIC"
        t.cell(0, 1).text = "Tugas"
        t.cell(1, 0).text = "STAFF-01"
        t.cell(1, 1).text = "Tahan cue"
        d.save(path)
        r = parse(path)
        self.assertEqual([b.kind for b in r.content_blocks], ["title", "heading", "paragraph", "list_item", "list_item", "table"])
        self.assertEqual(r.content_blocks[-1].data["rows"][1], ["STAFF-01", "Tahan cue"])
        self.assertEqual(r.content_blocks[2].source_locator["paragraph_position"], 3)
        self.assertEqual(r.content_blocks[-1].heading, "Procedure")

    def test_pdf_text_and_page_blocks(self):
        path = self.base / "sample.pdf"
        pdf_fixture(path)
        r = parse(path)
        self.assertEqual(r.content_blocks[0].kind, "heading")
        self.assertIn("Do not change STAFF-01 or 16:30.", r.content_blocks[-1].text)
        self.assertEqual(r.content_blocks[-1].source_locator["page_start"], 1)
        self.assertEqual(r.content_blocks[-1].source_locator["page_block"], 2)

    def test_pdf_without_text_fails_no_ocr(self):
        path = self.base / "blank.pdf"
        pdf_fixture(path, empty=True)
        with self.assertRaisesRegex(ParseError, "no text layer.*OCR disabled"):
            parse(path)

    def test_txt_lines_heading_paragraph_breaks_list(self):
        path = self.base / "sample.txt"
        path.write_text("Fixture\n\nProcedure\nJangan mengubah kode.\nBaris lanjut.\n\n- Tahan cue\n- Baca ulang\n", encoding="utf-8")
        blocks = parse(path).content_blocks
        self.assertEqual([b.kind for b in blocks], ["title", "heading", "paragraph", "list_item", "list_item"])
        self.assertEqual(blocks[2].source_locator, {"line_start": 4, "line_end": 5, "section": "Procedure"})
        self.assertIn("\n", blocks[2].text)

    def make_xlsx(self, path, formula=False):
        book = Workbook()
        s = book.active
        s.title = "Rundown"
        s.append(["Fixture"])
        s.append(["Waktu", "Aktivitas", "PIC"])
        s.append(["16:30", "Registrasi", "STAFF-08"])
        s.append(["17:00", "Tahan cue", "STAFF-01"])
        if formula:
            s["A3"] = "=1+1"
        book.save(path)

    def test_xlsx_header_row_and_sheet_context(self):
        path = self.base / "sample.xlsx"
        self.make_xlsx(path)
        blocks = parse(path).content_blocks
        row = blocks[2]
        self.assertEqual(row.kind, "table_row")
        self.assertEqual(row.source_locator["sheet"], "Rundown")
        self.assertEqual(row.source_locator["row_start"], 3)
        self.assertEqual(row.source_locator["header_row"], 2)
        self.assertEqual(row.data["headers"], ["Waktu", "Aktivitas", "PIC"])
        self.assertEqual(row.text, "Waktu: 16:30\nAktivitas: Registrasi\nPIC: STAFF-08")
        d = document(blocks)
        d.metadata["format"] = "XLSX"
        c = chunk_document(d, PILOT)[0]
        self.assertIn("Sheet: Rundown", c.text)
        self.assertIn("Row: 3", c.text)
        self.assertIn("Waktu: 16:30", c.text)
        self.assertEqual(c.source_locator["row_end"], 4)

    def test_xlsx_multiple_tables_and_merged_context(self):
        path = self.base / "sample.xlsx"
        b = Workbook()
        s = b.active
        s.append(["Section one"])
        s.merge_cells("A1:B1")
        s.append(["PIC", "Tugas"])
        s.append(["STAFF-01", "Tahan"])
        s.append([None, None])
        s.append(["Section two"])
        s.append(["Waktu", "Cue"])
        s.append(["16:30", "Standby"])
        b.save(path)
        blocks = parse(path).content_blocks
        self.assertEqual(blocks[0].source_locator["merged_ranges"], ["A1:B1"])
        rows = [v for v in blocks if v.kind == "table_row"]
        self.assertNotEqual(rows[0].source_locator["table"], rows[1].source_locator["table"])
        self.assertEqual(rows[1].data["headers"], ["Waktu", "Cue"])

    def test_missing_formula_cache_fails(self):
        path = self.base / "sample.xlsx"
        self.make_xlsx(path, formula=True)
        with self.assertRaisesRegex(ParseError, "Missing formula cache"):
            parse(path)

    def test_empty_files_and_unsupported_formats(self):
        for suffix in ["docx", "pdf", "txt", "xlsx"]:
            path = self.base / f"empty.{suffix}"
            path.touch()
            with self.assertRaisesRegex(ParseError, "Empty source file"):
                parse(path)
        path = self.base / "spaces.txt"
        path.write_text(" \n\n", encoding="utf-8")
        with self.assertRaisesRegex(ParseError, "No readable text"):
            parse(path)
        with self.assertRaisesRegex(ParseError, "Unsupported source format"):
            parse(self.base / "file.csv")


class CleaningChunkTests(unittest.TestCase):
    def test_markdown_report_stable_after_json_roundtrip(self):
        report = {"total_chunks": 1, "canonical_documents": 1}
        samples = [{"chunk_id": "WO-DOC-999-C001", "format": "TXT", "document_id": "WO-DOC-999", "heading": "Test",
                    "token_count": 1, "source_locator": {"line_start": 1, "line_end": 1}, "text_preview": "Test"}]
        boilerplate = {"policy": "exact", "patterns": {"version": 0, "draft": 0}}
        roundtrip = lambda value: json.loads(json.dumps(value, sort_keys=True))
        self.assertEqual(markdown_report(report, samples, boilerplate),
                         markdown_report(roundtrip(report), roundtrip(samples), roundtrip(boilerplate)))

    def test_boilerplate_exact_match_retains_operational_prose(self):
        texts = ["WO-DOC-999 | Versi 1.0 | Wedding Organizer X",
                 "Draft susunan peneliti; belum divalidasi praktisi. Skenario simulasi.",
                 "Wedding Organizer X tidak boleh mengubah STAFF-01 atau angka 120.",
                 "Rundown versi 1.0 jangan diganti tanpa EC."]
        d = clean(document([Block(str(i), "paragraph", text, {"paragraph_position": i + 1}) for i, text in enumerate(texts)]))
        self.assertEqual([b.retrievable for b in d.content_blocks], [False, False, True, True])
        self.assertEqual(d.content_blocks[0].raw_text, texts[0])
        self.assertIn("tidak boleh", chunk_document(d, PILOT)[0].text)

    def test_unicode_whitespace_safe(self):
        self.assertEqual(normalize(" e\u0301  STAFF-01\t16:30\n\n\nJangan  120 "), "é STAFF-01 16:30\n\nJangan 120")
        self.assertEqual(normalize("① ½ １２３"), "① ½ １２３")

    def test_small_sections_do_not_merge_or_overlap(self):
        blocks = [Block("1", "paragraph", "A" * 200, {"paragraph_position": 1}, "One"),
                  Block("2", "paragraph", "B" * 200, {"paragraph_position": 2}, "Two")]
        c = chunk_document(document(blocks), PILOT)
        self.assertEqual(len(c), 2)
        self.assertEqual([v.metadata["overlap_tokens"] for v in c], [0, 0])
        self.assertTrue(all("overlap_blocks" not in v.source_locator for v in c))

    def test_recursive_fallback_and_overlap_bound(self):
        text = " ".join(f"Kalimat {i} jangan ubah STAFF-01." for i in range(160))
        d = document([Block("1", "paragraph", text, {"line_start": 3, "line_end": 3}, "Procedure")])
        chunks = chunk_document(d, PILOT)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(any(c.metadata["overlap_tokens"] for c in chunks[1:]))
        self.assertTrue(all(c.token_count <= 500 for c in chunks))
        self.assertTrue(all(c.metadata["overlap_tokens"] <= 75 for c in chunks))
        for c in chunks:
            self.assertEqual(c.source_locator["line_start"], 3)
            self.assertEqual(c.source_locator["blocks"][0]["block_id"], "1")
        combined = " ".join(c.text for c in chunks)
        for i in range(160):
            self.assertIn(f"Kalimat {i} jangan ubah STAFF-01.", combined)

    def test_character_fallback_no_loss_and_stable_ids_hash(self):
        d = document([Block("1", "paragraph", "x" * 6200, {"paragraph_position": 1}, "Long")])
        first = chunk_document(d, PILOT)
        second = chunk_document(d, PILOT)
        self.assertEqual([asdict(c) for c in first], [asdict(c) for c in second])
        self.assertEqual(first[0].chunk_id, "WO-DOC-999-C001")
        for c in first:
            self.assertEqual(c.content_hash, hashlib.sha256(c.text.encode()).hexdigest())
            self.assertLessEqual(c.token_count, 500)

    def test_invalid_chunk_config_rejected(self):
        with self.assertRaises(ValueError):
            chunk_document(document([]), {**PILOT, "overlap_tokens": 500})


class CorpusIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = load_config(ROOT, ROOT / "experiments/configs/pilot.toml")
        cls.before = snapshot(ROOT / "data/source_corpus")
        cls.documents, cls.report, cls.audit = ingest(ROOT, cls.config)
        cls.chunks, cls.chunk_report = build_chunks(cls.documents, cls.config)

    def test_50_document_discovery_manifest_alignment(self):
        self.assertEqual(self.audit["source_documents"], 50)
        self.assertEqual(self.audit["errors"], [])
        self.assertEqual(self.audit["formats"], {"DOCX": 26, "PDF": 9, "TXT": 5, "XLSX": 10})

    def test_all_formats_parse_and_source_locators_preserved(self):
        self.assertEqual(self.report["successfully_parsed"], 50)
        self.assertEqual(self.report["failed_documents"], 0)
        self.assertEqual(self.report["formula_cells_read_from_cache"], 33)
        for c in self.chunks:
            self.assertTrue(c.source_file.startswith("data/source_corpus/documents/"))
            self.assertEqual(c.validation_status, "pending_practitioner_validation")
            self.assertTrue(c.source_locator["blocks"])
            fmt = c.metadata["format"]
            key = {"DOCX": "body_position", "PDF": "page_start", "TXT": "line_start", "XLSX": "row_start"}[fmt]
            self.assertIn(key, c.source_locator["blocks"][0])
            self.assertEqual(token_count(c.text), c.token_count)
            self.assertLessEqual(c.token_count, 500)

    def test_deterministic_documents_and_pipeline_repeatability(self):
        documents, report, _ = ingest(ROOT, self.config)
        chunks, chunk_report = build_chunks(documents, self.config)
        self.assertEqual(jsonl_text(documents), jsonl_text(self.documents))
        self.assertEqual(jsonl_text(chunks), jsonl_text(self.chunks))
        self.assertEqual(report, self.report)
        self.assertEqual(chunk_report, self.chunk_report)

    def test_canonical_serialization_roundtrip(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "documents.jsonl"
            path.write_text(jsonl_text(self.documents), encoding="utf-8")
            restored = read_documents(path)
            self.assertEqual(jsonl_text(restored), jsonl_text(self.documents))
            verify_canonical(ROOT, self.config, restored)

    def test_no_source_mutation(self):
        self.assertEqual(snapshot(ROOT / "data/source_corpus"), self.before)

    def test_source_output_paths_rejected(self):
        with self.assertRaisesRegex(ValueError, "unsafe output"):
            safe_write(ROOT, ROOT / "data/source_corpus/test.txt", "forbidden")

    def test_manifest_hash_mismatch_rejected_in_disposable_copy(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            shutil.copytree(ROOT / "data/source_corpus", root / "data/source_corpus")
            shutil.copy(ROOT / "README.md", root / "README.md")
            (root / "data/source_corpus/documents/WO-DOC-012.txt").write_text("Modified disposable copy", encoding="utf-8")
            _, report = audit(root, self.config)
            self.assertTrue(any("SHA-256 mismatch" in e for e in report["errors"]))
            with self.assertRaisesRegex(ValueError, "Source audit failed"):
                ingest(root, self.config)

    def test_no_administrative_status_in_retrievable_text(self):
        for c in self.chunks:
            self.assertNotIn("pending_practitioner_validation", c.text)
            self.assertNotIn("Draft susunan peneliti; belum divalidasi", c.text)
            self.assertNotIn("Wedding Organizer X | WO-DOC-", c.text)

    def test_xlsx_overlap_preserves_whole_rows(self):
        for c in self.chunks:
            if c.metadata["format"] != "XLSX" or not c.metadata["overlap_text"]:
                continue
            self.assertTrue(c.metadata["overlap_text"].startswith("Row: "))
            self.assertIn(c.source_locator["sheet"], c.text)
            for ref in c.source_locator["overlap_blocks"]:
                d = next(d for d in self.documents if d.document_id == c.document_id)
                b = next(b for b in d.content_blocks if b.block_id == ref["block_id"])
                self.assertIn(f"Row: {ref['row_start']}\n{b.text}", c.metadata["overlap_text"])

    def test_all_operational_block_text_covered_by_chunks(self):
        for d in self.documents:
            for b in d.content_blocks:
                if not b.retrievable or b.kind in {"heading", "title", "table_header"}:
                    continue
                original = f"Row: {b.source_locator['row_start']}\n{b.text}" if b.kind == "table_row" else b.text
                coverage = set()
                for c in self.chunks:
                    if c.document_id != d.document_id:
                        continue
                    for ref in c.source_locator["blocks"]:
                        if ref["block_id"] == b.block_id:
                            start, end = ref["serialized_char_start"], ref["serialized_char_end"]
                            self.assertIn(original[start:end], c.text)
                            coverage.update(range(start, end))
                self.assertTrue(all(i in coverage or char.isspace() for i, char in enumerate(original)), (d.document_id, b.block_id))


if __name__ == "__main__":
    unittest.main()
