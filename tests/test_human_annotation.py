"""Synthetic-only metric/persistence fixtures. Never write judgments into the research workspace."""
from copy import deepcopy
from hashlib import sha256
import json
import math
from pathlib import Path
import re
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import human_annotation as cli
from wo_retrieval import human_judgments as h
from wo_retrieval import human_metrics as m


def fixture(root):
    """Three synthetic query styles, twelve synthetic chunks; not WO research results."""
    root = Path(root); workspace = root / 'synthetic_human_workspace'; workspace.mkdir()
    rows = []
    styles = ('lexical_oriented', 'semantic_paraphrase', 'mixed_natural')
    queries = [{'query_id':f'TEST-Q{i}', 'text':f'SYNTHETIC TEST query {i}', 'style':style,
                'information_need_family_id':f'TEST-FAMILY{i}', 'purpose':'development'} for i,style in enumerate(styles)]
    for q in queries:
        for i in range(12):
            text = f'SYNTHETIC TEST FIXTURE chunk {i}; not a research judgment.'
            rows.append(dict(assessment_id=q['query_id']+f'-A{i:02d}', query_id=q['query_id'], chunk_id=f'TEST-C{i:02d}',
                document_id='TEST-DOC', query_text=q['text'], chunk_text=text, title='Synthetic test source',
                heading='Not research', source_locator={'test': True}, query_text_sha256=h.text_hash(q['text']),
                chunk_text_sha256=h.text_hash(text), snapshot_id='TEST-SNAPSHOT', snapshot_fingerprint='TEST-FINGERPRINT'))
    input_path=workspace/'annotation_input.jsonl';input_path.write_text(h.jsonl_text(rows), encoding='utf-8')
    frozen=root/'synthetic_frozen.txt'; frozen.write_text('synthetic frozen source',encoding='utf-8')
    manifest=dict(schema_version=1, package_id='SYNTHETIC-TEST-NOT-RESEARCH', scope='development',label=h.LABEL,
        assessor_id='TEST-ANNOTATOR',expert_validated=False,independently_validated=False,pair_count=len(rows),
        snapshot_id='TEST-SNAPSHOT',snapshot_fingerprint='TEST-FINGERPRINT',input_sha256=h.file_hash(input_path),
        frozen_source_hashes={'synthetic_frozen.txt':h.file_hash(frozen)},query_metadata=queries,
        protocol=h.PROTOCOL,protocol_sha256=h.digest(h.PROTOCOL))
    h.write_json(workspace/'annotation_manifest.json',manifest)
    h.write_json(workspace/'storage_manifest.json',dict(package_id=manifest['package_id'],input_sha256=manifest['input_sha256'],
        annotation_manifest_sha256=h.file_hash(workspace/'annotation_manifest.json'),assessor_id=manifest['assessor_id'],
        label=h.LABEL,scope='development',expert_validated=False,independently_validated=False))
    (workspace/'human_judgment_history.jsonl').write_text('',encoding='utf-8')
    package=h.load_package(root,workspace)
    rankings=[dict(query_id=q['query_id'],query_text=q['text'],method=method,results=[
        dict(chunk_id=f'TEST-C{i:02d}',document_id='TEST-DOC',rank=i+1,score=12.0-i,
             method_id=method,snapshot_id='TEST-SNAPSHOT') for i in range(12)]) for q in queries for method in h.METHODS]
    return package,rankings


def finish(package):
    for pair in package['rows']:
        index=int(pair['chunk_id'][-2:]);grade={1:1,2:2}.get(index,0)
        h.save_decision(package,pair['assessment_id'],'RESOLVED',grade,'SYNTHETIC TEST FIXTURE, NOT RESEARCH')
    return h.current_records(package)


class CalculationTests(unittest.TestCase):
    def test_precision_expected_values_and_short_windows(self):
        self.assertEqual(m.precision([0,1,2,0,0],5),2/5)
        self.assertEqual(m.precision([0,1,2],10),2/10)
        self.assertEqual(m.precision([],5),0)

    def test_reciprocal_rank_and_no_hit(self):
        self.assertEqual(m.reciprocal_rank([0,0,1,2]),1/3)
        self.assertEqual(m.reciprocal_rank([0]*10+[2]),0)
        self.assertEqual(m.reciprocal_rank([]),0)

    def test_graded_ndcg_mathematically_verified(self):
        expected=(1/math.log2(3)+3/math.log2(4))/(3+1/math.log2(3))
        self.assertAlmostEqual(m.ndcg([0,1,2,0,0],[2,1,0,0,0],5),expected,places=14)
        self.assertEqual(m.ndcg([2,1],[2,1],10),1)
        self.assertEqual(m.ndcg([],[2,1],10),0)

    def test_zero_ideal_is_blocked(self):
        with self.assertRaises(ValueError):m.ndcg([0],[0],5)

    def test_missing_unresolved_bool_and_invalid_grades(self):
        for grades in ([None],[True],[-1],[3],['1']):
            for func in (m.precision,m.reciprocal_rank):
                with self.assertRaises(ValueError):func(grades,10)
            with self.assertRaises(ValueError):m.ndcg(grades,[2],10)
        for k in (0,-1,True):
            with self.assertRaises(ValueError):m.precision([1],k)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.package,self.rankings=fixture(self.root)

    def test_initial_storage_has_no_prefilled_grades(self):
        records=h.current_records(self.package)
        self.assertTrue(all(r['grade'] is None and r['provenance'] is None for r in records))
        self.assertEqual(h.completeness(self.package,records)['missing'],36)

    def test_persistence_resume_and_revision_history(self):
        pair=self.package['rows'][0]
        h.save_decision(self.package,pair['assessment_id'],'RESOLVED',2,'test note',pair['chunk_text'])
        resumed=h.load_package(self.root,self.package['workspace'])
        record=h.current_records(resumed)[0]
        self.assertEqual(record['grade'],2);self.assertEqual(record['notes'],'test note')
        h.save_decision(resumed,pair['assessment_id'],'UNRESOLVED',None,'test ambiguity',expected_revision=1)
        events,latest=h.load_history(resumed)
        self.assertEqual(len(events),2);self.assertEqual(events[0]['grade'],2)
        self.assertIsNone(latest[pair['assessment_id']]['grade'])
        self.assertEqual(h.completeness(resumed,h.current_records(resumed))['unresolved'],1)

    def test_stale_revision_rejected(self):
        pair=self.package['rows'][0]
        h.save_decision(self.package,pair['assessment_id'],'RESOLVED',0)
        with self.assertRaises(ValueError):h.save_decision(self.package,pair['assessment_id'],'RESOLVED',2)
        self.assertEqual(len(h.load_history(self.package)[0]),1)

    def test_invalid_schema_and_fabricated_evidence(self):
        aid=self.package['rows'][0]['assessment_id']
        for status,grade in [('RESOLVED',True),('RESOLVED',None),('UNRESOLVED',0),('PENDING',0),('AI_SILVER',2)]:
            with self.assertRaises(ValueError):h.save_decision(self.package,aid,status,grade)
        with self.assertRaises(ValueError):h.save_decision(self.package,aid,'RESOLVED',1,evidence='fabricated quote')

    def test_atomic_save_failure_preserves_history(self):
        aid=self.package['rows'][0]['assessment_id']
        h.save_decision(self.package,aid,'RESOLVED',0)
        original=(self.package['workspace']/'human_judgment_history.jsonl').read_bytes()
        with patch('wo_retrieval.human_judgments.os.replace',side_effect=OSError('synthetic failure')):
            with self.assertRaises(OSError):h.save_decision(self.package,aid,'RESOLVED',1,expected_revision=1)
        self.assertEqual((self.package['workspace']/'human_judgment_history.jsonl').read_bytes(),original)
        self.assertFalse((self.package['workspace']/'.write.lock').exists())

    def test_snapshot_source_and_input_integrity(self):
        (self.root/'synthetic_frozen.txt').write_text('tampered',encoding='utf-8')
        with self.assertRaises(ValueError):h.load_package(self.root,self.package['workspace'])
        (self.root/'synthetic_frozen.txt').write_text('synthetic frozen source',encoding='utf-8')
        p=self.package['workspace']/'annotation_input.jsonl'
        p.write_text(p.read_text(encoding='utf-8')+'\n',encoding='utf-8')
        with self.assertRaises(ValueError):h.load_package(self.root,self.package['workspace'])

    def test_history_hash_tamper_rejected(self):
        h.save_decision(self.package,self.package['rows'][0]['assessment_id'],'RESOLVED',2)
        p=self.package['workspace']/'human_judgment_history.jsonl';events=h.read_jsonl(p)
        events[0]['grade']=0;p.write_text(h.jsonl_text(events),encoding='utf-8')
        with self.assertRaises(ValueError):h.load_history(self.package)

    def test_duplicate_pairs_and_records_rejected(self):
        records=h.current_records(self.package)
        with self.assertRaises(ValueError):m.evaluate(self.package,records+[records[0]],self.rankings)
        with self.assertRaises(ValueError):h.unique([self.package['rows'][0]]*2,('query_id','chunk_id'))
        with self.assertRaises(ValueError):m.evaluate(self.package,records,self.rankings+[self.rankings[0]])

    def test_gates_for_missing_and_unresolved(self):
        result=m.evaluate(self.package,h.current_records(self.package),self.rankings)
        self.assertEqual(result['status'],'BLOCKED_INCOMPLETE_ASSESSMENTS')
        self.assertTrue(all(r[metric] is None for r in result['macro'] for metric in m.METRICS))
        records=finish(self.package);aid=self.package['rows'][0]['assessment_id']
        h.save_decision(self.package,aid,'UNRESOLVED',None,expected_revision=1)
        result=m.evaluate(self.package,h.current_records(self.package),self.rankings)
        self.assertEqual(result['completeness']['unresolved'],1)
        self.assertTrue(all(r['evaluated_queries']==0 for r in result['macro']))

    def test_metric_values_macro_styles_and_descriptive_comparisons(self):
        records=finish(self.package);result=m.evaluate(self.package,records,self.rankings)
        expected=(1/math.log2(3)+3/2)/(3+1/math.log2(3))
        self.assertEqual(result['status'],'VALID_COMPLETED_DEVELOPMENT_POOLED')
        self.assertEqual(len(result['per_query']),12);self.assertEqual(len(result['by_style']),12)
        for r in result['macro']:
            self.assertEqual(r['Precision@5'],.4);self.assertEqual(r['Precision@10'],.2)
            self.assertEqual(r['MRR@10'],.5);self.assertAlmostEqual(r['nDCG@10 (pooled)'],expected)
        self.assertTrue(all(r['ties']==3 for r in result['method_comparisons']))
        self.assertEqual(result,m.evaluate(self.package,records,list(reversed(self.rankings))))

    def test_nonzero_method_comparisons(self):
        records=finish(self.package)
        ranking=self.rankings[0]['results'];ranking[0],ranking[2]=ranking[2],ranking[0]
        for i,r in enumerate(ranking):r['rank']=i+1;r['score']=12.0-i
        result=m.evaluate(self.package,records,self.rankings)
        contrast=next(r for r in result['method_comparisons'] if r['method_a']=='R1' and r['method_b']=='R2' and r['metric']=='MRR@10')
        self.assertEqual(contrast['wins_a'],1);self.assertEqual(contrast['ties'],2)
        self.assertAlmostEqual(contrast['mean_delta_a_minus_b'],1/6)

    def test_zero_relevant_pool_blocks_no_answer_inference(self):
        for row in self.package['rows']:h.save_decision(self.package,row['assessment_id'],'RESOLVED',0)
        result=m.evaluate(self.package,h.current_records(self.package),self.rankings)
        self.assertEqual(result['status'],'BLOCKED_NO_KNOWN_RELEVANCE')
        self.assertEqual(len(result['zero_ideal_query_ids']),3)

    def test_tie_order_and_duplicate_rank_candidates(self):
        records=finish(self.package)
        for row in self.rankings[0]['results']:row['score']=1.0
        self.assertEqual(m.evaluate(self.package,records,self.rankings)['status'],'VALID_COMPLETED_DEVELOPMENT_POOLED')
        self.rankings[0]['results'][0],self.rankings[0]['results'][1]=self.rankings[0]['results'][1],self.rankings[0]['results'][0]
        with self.assertRaises(ValueError):m.evaluate(self.package,records,self.rankings)

    def test_duplicate_candidates_and_altered_document_binding(self):
        records=finish(self.package);bad=deepcopy(self.rankings)
        bad[0]['results'].append(bad[0]['results'][0])
        with self.assertRaises(ValueError):m.evaluate(self.package,records,bad)
        bad=deepcopy(self.rankings);bad[0]['results'][0]['document_id']='wrong document'
        with self.assertRaises(ValueError):m.evaluate(self.package,records,bad)

    def test_short_returned_list_is_not_missing_judgment(self):
        records=finish(self.package)
        for ranking in self.rankings:ranking['results']=ranking['results'][:3]
        result=m.evaluate(self.package,records,self.rankings)
        self.assertEqual(result['status'],'VALID_COMPLETED_DEVELOPMENT_POOLED')
        self.assertEqual(result['macro'][0]['Precision@10'],.2)
        self.assertEqual(result['window_coverage'][1]['missing_slots'],7)

    def test_binary_threshold_frozen_and_ai_provenance_rejected(self):
        with self.assertRaises(ValueError):m.precision([1,2],5,2)
        records=finish(self.package);records[0]['provenance']['label']='AI SILVER — PROVISIONAL'
        with self.assertRaises(ValueError):m.evaluate(self.package,records,self.rankings)

    def test_main_scope_is_not_development(self):
        package=deepcopy(self.package);package['manifest']['scope']='evaluation'
        with self.assertRaises(ValueError):m.evaluate(package,h.current_records(self.package),self.rankings)

    def test_report_nulls_and_word_completion_gate(self):
        result=m.evaluate(self.package,h.current_records(self.package),self.rankings)
        out=self.root/'blocked_report';m.publish_report(result,out)
        self.assertIn('NOT COMPUTED',(out/'summary.md').read_text(encoding='utf-8'))
        with self.assertRaises(ValueError):m.write_word(result,self.root/'blocked.docx')
        self.assertFalse((self.root/'blocked.docx').exists())
        with self.assertRaises(FileExistsError):m.publish_report(result,out)

    def test_word_format_editable_and_fixture_is_labeled(self):
        from docx import Document
        result=m.evaluate(self.package,finish(self.package),self.rankings)
        path=self.root/'synthetic_not_research.docx';m.write_word(result,path,synthetic=True)
        doc=Document(path)
        self.assertIn('SYNTHETIC TEST FIXTURE',doc.paragraphs[0].text)
        self.assertEqual(len(doc.tables),7)
        self.assertAlmostEqual(doc.sections[0].page_width.cm,21,places=2)
        self.assertIn('SYNTHETIC TEST FIXTURE',doc.sections[0].header.paragraphs[0].text)

    def test_export_preserves_unknown_status_and_provenance(self):
        h.save_decision(self.package,self.package['rows'][0]['assessment_id'],'UNRESOLVED',None)
        out=self.root/'synthetic_export';cli.export(self.package,out)
        records=h.read_jsonl(out/'human_judgments.jsonl')
        self.assertEqual(records[0]['status'],'UNRESOLVED');self.assertIsNone(records[0]['grade'])
        self.assertEqual((out/'human_resolved_qrels.tsv').read_text(), '')
        self.assertEqual(h.read_json(out/'export_manifest.json')['completeness']['missing'],35)

    def test_http_blinding_autosave_and_local_access(self):
        server=cli.make_server(self.package,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        self.addCleanup(lambda:(server.shutdown(),server.server_close(),thread.join()))
        url=f'http://127.0.0.1:{server.server_port}'
        html=urlopen(url).read().decode('utf-8');token=json.loads(re.search(r'const token=(.*?);',html).group(1))
        state=json.loads(urlopen(url+'/api/state').read())
        self.assertEqual(set(state['pair']),h.PAIR_FIELDS)
        for forbidden in ('method','score','rank','silver_grade','existing_grade','query_style'):
            self.assertNotIn(forbidden,state['pair'])
        payload=dict(token=token,assessment_id=state['pair']['assessment_id'],expected_revision=0,
                     status='RESOLVED',grade=1,notes='SYNTHETIC HTTP TEST',evidence='')
        response=json.loads(urlopen(Request(url+'/api/save',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'})).read())
        self.assertEqual(response['progress']['resolved'],1)
        self.assertEqual(json.loads(urlopen(url+'/api/state?index=0').read())['judgment']['grade'],1)
        for headers in ({'Origin':'https://external.invalid'},{'Host':'external.invalid'}):
            with self.assertRaises(HTTPError) as error:urlopen(Request(url+'/api/state',headers=headers))
            self.assertEqual(error.exception.code,403)
        payload['token']='invalid'
        with self.assertRaises(HTTPError) as error:urlopen(Request(url+'/api/save',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'}))
        self.assertEqual(error.exception.code,403)
        with self.assertRaises(HTTPError):urlopen(url+'/annotation_manifest.json')


class SplitTests(unittest.TestCase):
    def test_family_and_text_separation_and_semantic_review(self):
        dev=[dict(query_id='DEV',purpose='development',text='Same Need',information_need_family_id='F1')]
        ev=[dict(query_id='MAIN',purpose='evaluation',text='Different need',information_need_family_id='F2')]
        h.validate_split(dev,ev,True)
        with self.assertRaises(ValueError):h.validate_split(dev,ev)
        for patch_value in ({'query_id':'DEV'},{'text':'  SAME   need '},{'information_need_family_id':'F1'}):
            with self.assertRaises(ValueError):h.validate_split(dev,[{**ev[0],**patch_value}],True)


class PublishedPackageTests(unittest.TestCase):
    def test_actual_package_is_blinded_and_bound_to_all_654_pairs(self):
        root=Path(__file__).resolve().parents[1];workspace=root/h.WORKSPACE
        if not workspace.exists():self.skipTest('Package not yet prepared')
        p=h.load_package(root,workspace)
        self.assertEqual(len(p['rows']),654);self.assertEqual(len(p['manifest']['query_metadata']),18)
        self.assertTrue(all(r['status'] in ('PENDING','RESOLVED','UNRESOLVED') for r in h.current_records(p)))
        self.assertTrue(all(set(r)==h.PAIR_FIELDS for r in p['rows']))
        self.assertEqual(p['manifest']['seed'],17042)
        sorted_rows=sorted(p['rows'],key=lambda r:(r['query_id'],r['chunk_id']))
        import random
        random.Random(17042).shuffle(sorted_rows)
        self.assertEqual(sorted_rows,p['rows'])


if __name__=='__main__':unittest.main()
