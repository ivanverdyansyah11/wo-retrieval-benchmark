"""Human-only pooled development metrics; no retrieval or automatic judging."""
from itertools import combinations
import csv
import math
from pathlib import Path
from statistics import mean

from .human_judgments import (LABEL, METHODS, PROTOCOL, binding, completeness, digest,
                              file_hash, unique, validate_decision, write_json)

METRICS = ('Precision@5', 'Precision@10', 'MRR@10', 'nDCG@5 (pooled)', 'nDCG@10 (pooled)')


def precision(grades, k, threshold=1):
    validate_grades(grades)
    validate_k(k)
    if threshold != 1:
        raise ValueError('This frozen protocol uses binary threshold >=1')
    return sum(g >= threshold for g in grades[:k]) / k


def reciprocal_rank(grades, k=10, threshold=1):
    validate_grades(grades); validate_k(k)
    if threshold != 1:
        raise ValueError('This frozen protocol uses binary threshold >=1')
    return next((1 / rank for rank, grade in enumerate(grades[:k], 1) if grade >= threshold), 0.0)


def validate_grades(grades):
    if any(type(g) is not int or g not in (0, 1, 2) for g in grades):
        raise ValueError('Only resolved integer grades accepted; missing/null is not zero')


def validate_k(k):
    if type(k) is not int or k < 1:
        raise ValueError('k must be a positive integer')


def dcg(grades, k):
    validate_grades(grades); validate_k(k)
    return sum((2**grade - 1) / math.log2(rank + 1) for rank, grade in enumerate(grades[:k], 1))


def ndcg(grades, ideal_grades, k):
    ideal = dcg(sorted(ideal_grades, reverse=True), k)
    actual = dcg(grades, k)
    if ideal == 0:
        raise ValueError('Zero pooled IDCG: review answerability; do not infer no relevant corpus chunks')
    if actual > ideal + 1e-12:
        raise ValueError('Ideal scope is inconsistent with ranked judgments')
    return actual / ideal


def evaluate(package, records, rankings):
    records = sorted(records, key=lambda r: r['assessment_id'])
    rankings = sorted(rankings, key=lambda r: (r['query_id'], r['method']))
    manifest = package['manifest']
    if manifest['scope'] != 'development' or manifest['protocol'] != PROTOCOL:
        raise ValueError('Main evaluation is not authorized by the development package')
    recordmap = unique(records, ('assessment_id',))
    unique([r['binding'] for r in records], ('query_id', 'chunk_id'))
    if set(k[0] for k in recordmap) != set(package['pairs']):
        raise ValueError('Missing/extra human records; export PENDING explicitly')
    for record in records:
        pair = package['pairs'][record['assessment_id']]
        if record['binding'] != binding(package, pair):
            raise ValueError('Human query/chunk/snapshot binding mismatch')
        validate_decision(record['status'], record['grade'], record['notes'], record['evidence'], pair)
        if record['status'] != 'PENDING':
            provenance = record['provenance']
            if (not provenance or provenance['label'] != LABEL
                    or provenance['assessor_id'] != manifest['assessor_id']
                    or provenance['scope'] != manifest['scope']
                    or provenance['protocol_sha256'] != manifest['protocol_sha256']
                    or provenance['expert_validated'] is not False or provenance['independently_validated'] is not False):
                raise ValueError('Nonhuman/unbound judgment rejected')
    qmap = {q['query_id']: q for q in manifest['query_metadata']}
    runs = unique(rankings, ('query_id', 'method'))
    if set(runs) != {(q, m) for q in qmap for m in METHODS}:
        raise ValueError('Compared methods require identical query scope')
    grades = {(r['binding']['query_id'], r['binding']['chunk_id']): r for r in records}
    windows, missing_records = [], []
    for (qid, method), run in sorted(runs.items()):
        if run['query_text'] != qmap[qid]['text']:
            raise ValueError('Ranking query text mismatch')
        unique(run['results'], ('chunk_id',))
        if run['results'] != sorted(run['results'], key=lambda r: (-r['score'], r['chunk_id'])):
            raise ValueError('Ranking score/tie order changed')
        for i, row in enumerate(run['results'], 1):
            record = grades.get((qid, row['chunk_id']))
            if (type(row['rank']) is not int or row['rank'] != i
                    or type(row['score']) not in (int, float) or not math.isfinite(row['score'])
                    or row['snapshot_id'] != manifest['snapshot_id'] or row['method_id'] != method
                    or record and row['document_id'] != record['binding']['document_id']):
                raise ValueError('Ranking metadata invalid')
        for k in (5, 10):
            window = run['results'][:k]
            missing = unresolved = resolved = 0
            for row in window:
                record = grades.get((qid, row['chunk_id']))
                if record is None or record['status'] == 'PENDING':
                    missing += 1
                elif record['status'] == 'UNRESOLVED':
                    unresolved += 1
                else:
                    resolved += 1
                if record is None or record['status'] != 'RESOLVED':
                    missing_records.append(dict(query_id=qid, method=method, k=k, chunk_id=row['chunk_id'],
                        status=record['status'] if record else 'OUTSIDE_ASSESSMENT_SCOPE'))
            windows.append(dict(query_id=qid, method=method, k=k, returned=len(window), missing_slots=k-len(window),
                resolved=resolved, missing=missing, unresolved=unresolved,
                resolved_fraction=resolved/len(window) if window else None,
                complete=missing == unresolved == 0))
    coverage = completeness(package, records)
    status = 'BLOCKED_INCOMPLETE_ASSESSMENTS'
    zero_ideal = []
    if coverage['status'] == 'COMPLETE_RESOLVED' and all(w['complete'] for w in windows):
        zero_ideal = [qid for qid in sorted(qmap) if not any(r['grade'] >= 1 for r in records if r['binding']['query_id'] == qid)]
        status = 'BLOCKED_NO_KNOWN_RELEVANCE' if zero_ideal else 'VALID_COMPLETED_DEVELOPMENT_POOLED'
    per_query = []
    for qid in sorted(qmap):
        ideal = [r['grade'] for r in records if r['binding']['query_id'] == qid]
        for method in METHODS:
            values = {metric: None for metric in METRICS}
            if status == 'VALID_COMPLETED_DEVELOPMENT_POOLED':
                ranked = [grades[qid, row['chunk_id']]['grade'] for row in runs[qid, method]['results'][:10]]
                values = {'Precision@5': precision(ranked, 5), 'Precision@10': precision(ranked, 10),
                          'MRR@10': reciprocal_rank(ranked), 'nDCG@5 (pooled)': ndcg(ranked, ideal, 5),
                          'nDCG@10 (pooled)': ndcg(ranked, ideal, 10)}
            per_query.append(dict(query_id=qid, query_text=qmap[qid]['text'], style=qmap[qid]['style'], method=method, status=status, **values))
    macro, by_style = [], []
    for method in METHODS:
        for style in [None] + sorted({q['style'] for q in qmap.values()}):
            subset = [r for r in per_query if r['method'] == method and (style is None or r['style'] == style)]
            summary = dict(method=method, style=style or 'all', requested_queries=len(subset),
                evaluated_queries=len(subset) if status == 'VALID_COMPLETED_DEVELOPMENT_POOLED' else 0,
                status=status, **{m: mean(r[m] for r in subset) if status == 'VALID_COMPLETED_DEVELOPMENT_POOLED' else None for m in METRICS})
            (macro if style is None else by_style).append(summary)
    contrasts = []
    if status == 'VALID_COMPLETED_DEVELOPMENT_POOLED':
        for a, b in combinations(METHODS, 2):
            for metric in METRICS:
                deltas = [next(r[metric] for r in per_query if r['query_id'] == q and r['method'] == a)
                          - next(r[metric] for r in per_query if r['query_id'] == q and r['method'] == b) for q in sorted(qmap)]
                contrasts.append(dict(method_a=a, method_b=b, metric=metric, mean_delta_a_minus_b=mean(deltas),
                    wins_a=sum(d > 1e-12 for d in deltas), ties=sum(abs(d) <= 1e-12 for d in deltas),
                    losses_a=sum(d < -1e-12 for d in deltas), statistical_significance='NOT TESTED'))
    return dict(schema_version=1, status=status, calculator_readiness='TESTED_SEPARATELY_WITH_SYNTHETIC_FIXTURES',
        scope='development pooled evaluation', main_evaluation='NOT RUN / NOT ENABLED',
        judgment_label=LABEL, expert_validated=False, independently_validated=False,
        protocol=PROTOCOL, protocol_sha256=manifest['protocol_sha256'],
        package_id=manifest['package_id'], input_sha256=manifest['input_sha256'],
        snapshot_id=manifest['snapshot_id'], snapshot_fingerprint=manifest['snapshot_fingerprint'],
        human_records_sha256=digest(records), ranking_records_sha256=digest(rankings),
        completeness=coverage, window_coverage=windows, missing_unresolved_windows=missing_records,
        pending_unresolved_pairs=[{'assessment_id': r['assessment_id'], **r['binding'], 'status': r['status']}
                                  for r in records if r['status'] != 'RESOLVED'],
        zero_ideal_query_ids=zero_ideal, macro=macro, per_query=per_query, by_style=by_style,
        method_comparisons=contrasts, recall='NOT COMPUTED; no corpus-wide relevant denominator',
        warning='Single researcher; synthetic corpus; development queries; pooled ideal denominator. No final effectiveness or significance claim.')


def write_csv(path, rows):
    if not rows:
        return
    with Path(path).open('x', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)


def publish_report(result, out):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=False)
    write_json(out / 'evaluation.json', result)
    for name in ('macro', 'per_query', 'by_style', 'window_coverage', 'method_comparisons'):
        write_csv(out / (name + '.csv'), result[name])
    lines = ['# Pooled development evaluation', '', result['status'], '', result['warning'], '',
             'Main evaluation: NOT RUN. Null cells mean NOT COMPUTED, not zero.', '',
             '| Method | Precision@5 | Precision@10 | MRR@10 | nDCG@10 (pooled) |', '|---|---:|---:|---:|---:|']
    for row in result['macro']:
        vals = [row[m] for m in ('Precision@5','Precision@10','MRR@10','nDCG@10 (pooled)')]
        lines.append('| ' + row['method'] + ' | ' + ' | '.join('NOT COMPUTED' if v is None else f'{v:.6f}' for v in vals) + ' |')
    lines += ['', 'nDCG@5 tersedia pada JSON/CSV. Ideal ranking mencakup seluruh human-resolved pool per query.',
              'Precision/MRR: grade >=1. nDCG: gain 2^grade-1; discount log2(rank+1).',
              f"Completion: {result['completeness']['resolved']}/{result['completeness']['total']} RESOLVED; "
              f"{result['completeness']['missing']} missing; {result['completeness']['unresolved']} UNRESOLVED."]
    (out / 'summary.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    write_json(out / 'report_manifest.json', dict(status=result['status'], scope=result['scope'],
        human_records_sha256=result['human_records_sha256'], ranking_records_sha256=result['ranking_records_sha256'],
        protocol_sha256=result['protocol_sha256'], outputs={p.name:file_hash(p) for p in sorted(out.iterdir())}))


def write_word(result, path, synthetic=False):
    """Generate Word only after a complete result, with explicit fixture labeling in tests."""
    if result['status'] != 'VALID_COMPLETED_DEVELOPMENT_POOLED':
        raise ValueError('Word metrics report requires complete human-resolved pooled judgments')
    path = Path(path)
    if path.exists():
        raise ValueError('Word output exists; refusing overwrite')
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(); section = doc.sections[0]
    section.page_width=Cm(21); section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(2.5)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name]; style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
        for border in style.element.findall('.//'+qn('w:pBdr')):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11); doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Halaman ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    if synthetic:
        section.header.paragraphs[0].text='SYNTHETIC TEST FIXTURE — NOT RESEARCH RESULTS'
        doc.add_paragraph('SYNTHETIC TEST FIXTURE — NOT RESEARCH RESULTS', 'Title')
    doc.add_heading('Hasil Evaluasi Development Retrieval', 0)
    doc.add_paragraph('Pooled Development Evaluation — Not Final Thesis Evaluation')
    doc.add_paragraph('Laporan ini membandingkan empat metode pada 18 development queries dengan penilaian relevansi peneliti. Hasil berlaku untuk cakupan pilot dan pool yang dinilai, belum merupakan hasil evaluasi utama skripsi.')
    doc.add_paragraph('R1 adalah TF-IDF; R2 adalah BM25; R3 adalah BGE-M3 Dense; R4 adalah Hybrid BM25 dan Dense dengan Reciprocal Rank Fusion.')
    doc.add_paragraph(LABEL + '. Penilai: Muhammad Ivan Verdyansyah. Belum expert-validated atau independently validated.')
    doc.add_paragraph('Precision/MRR: grade ≥1. nDCG: gain 2^grade−1, discount log2(rank+1), ideal ranking dari seluruh union Top-20 pool per query. Recall tidak dihitung. Di luar pool tetap unknown.')
    doc.add_paragraph('Snapshot: '+result['snapshot_id'])
    doc.add_heading('Cara membaca metrik',1)
    doc.add_paragraph('Precision menunjukkan proporsi hasil relevan pada lima atau sepuluh posisi pertama. MRR menunjukkan seberapa awal hasil relevan pertama ditemukan dalam sepuluh posisi. nDCG memberi bobot lebih besar pada hasil sangat relevan yang muncul lebih awal, dibanding urutan ideal dalam pool yang sama. Semua metrik bernilai 0 hingga 1; lebih tinggi berarti lebih baik menurut metrik tersebut pada scope ini. Macro-average memberi bobot yang sama pada setiap query.')
    def add_table(headers, rows, widths):
        table=doc.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        borders=OxmlElement('w:tblBorders')
        for name in ('top','left','bottom','right','insideH','insideV'):
            edge=OxmlElement('w:'+name); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
        table._tbl.tblPr.append(borders)
        for col,width in zip(table.columns,widths): col.width=Cm(width)
        repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index,row in enumerate([headers]+rows):
            cells=table.rows[0].cells if index==0 else table.add_row().cells
            table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for i,text in enumerate(row):
                cell=cells[i];cell.width=Cm(widths[i]);cell.text=str(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins=OxmlElement('w:tcMar')
                for edge in ('top','bottom','left','right'):
                    item=OxmlElement('w:'+edge);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');margins.append(item)
                cell._tc.get_or_add_tcPr().append(margins)
                if index==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF3');cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(0)
                    paragraph.paragraph_format.keep_with_next=index<len(rows)
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:run.font.name='Arial';run.font.size=Pt(9.5);run.bold=index==0
        doc.add_paragraph('')
    metrics=METRICS
    doc.add_paragraph('Angka tabel dibulatkan menjadi empat angka desimal. Data perhitungan dengan presisi penuh tersedia pada berkas hasil pendamping.')
    format_values=lambda row:[f'{row[m]:.4f}' for m in metrics]
    doc.add_heading('Macro-average seluruh development queries',1)
    add_table(['Method',*metrics],[[r['method'],*format_values(r)] for r in result['macro']],[2,2.7,2.7,2,3.3,3.3])
    doc.add_heading('Hasil menurut query style',1)
    for style in sorted({r['style'] for r in result['by_style']}):
        doc.add_heading(style,2)
        add_table(['Method',*metrics],[[r['method'],*format_values(r)] for r in result['by_style'] if r['style']==style],[2,2.7,2.7,2,3.3,3.3])
    doc.add_heading('Hasil per query',1)
    for qid in sorted({r['query_id'] for r in result['per_query']}):
        doc.add_heading(qid,2)
        query_paragraph=doc.add_paragraph(next(r['query_text'] for r in result['per_query'] if r['query_id']==qid))
        query_paragraph.paragraph_format.keep_with_next=True
        add_table(['Method',*metrics],[[r['method'],*format_values(r)] for r in result['per_query'] if r['query_id']==qid],[2,2.7,2.7,2,3.3,3.3])
    doc.add_heading('Kelengkapan dan batas interpretasi',1)
    c=result['completeness'];doc.add_paragraph(f"{c['resolved']}/{c['total']} RESOLVED; missing {c['missing']}; UNRESOLVED {c['unresolved']}.")
    doc.add_paragraph('Main evaluation belum dijalankan. Tidak ada kesimpulan metode terbaik atau signifikansi statistik. Hasil ini menggunakan corpus sintetis, development queries, dan satu penilai peneliti. nDCG pooled dapat berbeda dari nDCG dengan exhaustive corpus judgments.')
    doc.add_paragraph('Descriptive paired method differences tersedia pada JSON/CSV yang menyertai laporan; bukan pengujian signifikansi.')
    doc.add_heading('Interpretasi hasil development',1)
    for metric in ('Precision@10','MRR@10','nDCG@10 (pooled)'):
        best=max(r[metric] for r in result['macro'])
        leaders=', '.join(r['method'] for r in result['macro'] if abs(r[metric]-best)<1e-12)
        doc.add_paragraph(f'Pada {metric}, macro-average tertinggi dalam development set ini adalah {leaders} dengan nilai {best:.6f}. Perbedaan ini bersifat deskriptif; bukan bukti keunggulan statistik atau generalisasi pada evaluation queries independen.')
    doc.add_paragraph('Langkah berikutnya adalah menyepakati evaluation queries independen dan protokol bersama pembimbing, membekukan rancangan, lalu menilai relevansinya dalam workspace terpisah. Rencana 75 evaluation queries belum difinalisasi. Review praktisi WO atau second human assessor dapat ditambahkan sebagai tahap validasi tersendiri.')
    doc.core_properties.author='Muhammad Ivan Verdyansyah';doc.core_properties.title='Hasil Evaluasi Development Retrieval'
    if synthetic:doc.core_properties.subject='Synthetic test fixture; not research results'
    doc.save(path)
